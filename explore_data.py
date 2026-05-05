import pandas as pd
import openpyxl
import os

dataset_dir = 'dataset/signaux'
files = [f for f in os.listdir(dataset_dir) if f.endswith('.xlsx')]

for file in sorted(files):
    filepath = os.path.join(dataset_dir, file)
    print(f'\n{"="*70}')
    print(f'FILE: {file}')
    print("="*70)
    
    try:
        wb = openpyxl.load_workbook(filepath)
        print(f'Sheets: {wb.sheetnames}')
        
        for sheet_name in wb.sheetnames:
            df = pd.read_excel(filepath, sheet_name=sheet_name)
            print(f'\n--- Sheet "{sheet_name}": Shape {df.shape} ---')
            print(f'Columns: {list(df.columns)}')
            print(f'Data types:\n{df.dtypes}')
            print(f'\nFirst rows:')
            print(df.head(5))
            print(f'\nLast rows:')
            print(df.tail(3))
    except Exception as e:
        print(f'ERROR: {e}')
